# -*- coding: utf-8 -*-
"""自动调整执行申请书（01）行距，消除正文末页大块空白并保留独立落款页。

用法：python fit_application_layout.py "<01执行申请书.docx>"
"""

import os
import sys
import tempfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pdfplumber
import win32com.client


def set_line_spacing(doc, twips):
    for p in doc.paragraphs:
        ppr = p._p.get_or_add_pPr()
        spacing = ppr.find(qn("w:spacing"))
        if spacing is None:
            spacing = OxmlElement("w:spacing")
            ppr.append(spacing)
        spacing.set(qn("w:line"), str(twips))
        spacing.set(qn("w:lineRule"), "exact")


def body_bottom(page):
    cutoff = page.height * 0.94  # 排除页脚页码
    return max((c["bottom"] for c in page.chars if c["bottom"] < cutoff), default=0.0)


def main():
    if len(sys.argv) != 2:
        print(__doc__)
        return 1
    docx_path = os.path.abspath(sys.argv[1])
    if not os.path.isfile(docx_path):
        print("找不到文件：", docx_path)
        return 1

    app = win32com.client.Dispatch("KWPS.Application")
    app.Visible = False
    app.DisplayAlerts = False
    candidates = []
    try:
        for line in range(480, 721, 10):
            doc = Document(docx_path)
            set_line_spacing(doc, line)
            tmp_docx = os.path.join(
                tempfile.gettempdir(), "_fit_01_%d_%d.docx" % (os.getpid(), line)
            )
            tmp_pdf = os.path.join(
                tempfile.gettempdir(), "_fit_01_%d_%d.pdf" % (os.getpid(), line)
            )
            doc.save(tmp_docx)
            word = None
            try:
                word = app.Documents.Open(tmp_docx, False, True)
                word.SaveAs(tmp_pdf, 17)
            finally:
                if word is not None:
                    word.Close(0)
            try:
                with pdfplumber.open(tmp_pdf) as pdf:
                    pages = pdf.pages
                    last_text = (pages[-1].extract_text() or "").strip()
                    if len(pages) > 1 and last_text.startswith("此致"):
                        bottoms = [body_bottom(p) for p in pages[:-1]]
                        if bottoms:
                            candidates.append((len(pages), bottoms[-1], line))
            finally:
                for path in (tmp_docx, tmp_pdf):
                    if os.path.isfile(path):
                        os.remove(path)
    finally:
        app.Quit()

    if not candidates:
        print("未找到“正文页＋独立落款末页”的可用行距")
        return 1

    min_pages = min(item[0] for item in candidates)
    same_pages = [item for item in candidates if item[0] == min_pages]
    _, bottom, selected = max(same_pages, key=lambda item: item[1])

    final_doc = Document(docx_path)
    set_line_spacing(final_doc, selected)
    tmp_final = docx_path + ".tmp"
    final_doc.save(tmp_final)
    os.replace(tmp_final, docx_path)
    print(
        "selected_line_twips=%d pages=%d last_body_bottom=%.1f"
        % (selected, min_pages, bottom)
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
