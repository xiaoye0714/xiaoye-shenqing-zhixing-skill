# -*- coding: utf-8 -*-
"""执行材料一次性复核：批量转 PDF，程序化检查，清理临时文件。

用法：
    python review_case.py "<案件文件夹>" [旧案关键词 ...]

做四件事（全程一条命令，PDF/PNG 只落临时目录且收尾自动删除）：
1. 用 WPS COM（KWPS.Application，隐藏窗口）把案件文件夹里的 docx 一次性全部导成 PDF；
2. 用 pdfplumber 逐份检查：页数、落款（此致/申请执行人/日期）是否同页、正文有没有被裁行
   （docx 文本是否都能在 PDF 文本里找到）、是否命中旧案关键词；
3. 打印紧凑结果表；只有被标 [!] 的文件才需要用 pdftoppm 渲染成 PNG 人工看；
4. 退出前删掉临时 PDF 目录。

要求目标 docx 没被 WPS 占用；被占用会打印是哪几份，先关掉再跑（不要一份份试）。
本脚本只关闭自己打开的那几份，绝不调用 app.Quit()——否则会把用户自己打开的 WPS 文档一并关掉。
"""

import glob
import os
import re
import shutil
import subprocess
import sys
import tempfile


def norm(s):
    return re.sub(r"\s+", "", s or "")


def convert_all(folder, outdir):
    import win32com.client
    app = win32com.client.Dispatch("KWPS.Application")
    try:
        app.Visible = False
        app.DisplayAlerts = False
    except Exception:
        pass
    made = []
    files = [f for f in sorted(glob.glob(os.path.join(folder, "*.docx")))
             if not os.path.basename(f).startswith("~$")]
    for f in files:
        name = os.path.basename(f)
        pdf = os.path.join(outdir, os.path.splitext(name)[0] + ".pdf")
        doc = None
        try:
            doc = app.Documents.Open(f, False, True)
            doc.SaveAs(pdf, 17)
            made.append((name, f, pdf))
        except Exception as exc:
            print("  [转换失败] %s -> %s" % (name, exc))
        finally:
            if doc is not None:
                try:
                    doc.Close(0)
                except Exception:
                    pass
    # 绝不 app.Quit()：COM 会挂到用户正在用的 WPS 实例上，Quit 会把用户打开的文档一起关掉。
    # 只 Close 本脚本打开的那几份即可。
    return made


def page_texts(pdf):
    import pdfplumber
    with pdfplumber.open(pdf) as doc:
        return [(pg.extract_text() or "") for pg in doc.pages]


def doc_texts(docx):
    from docx import Document
    d = Document(docx)
    out = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    out.append(p.text)
    return out


def shingle_ratio(t, flat, k=6):
    t = norm(t)
    if len(t) < k:
        return 1.0
    grams = [t[i:i + k] for i in range(0, len(t) - k + 1, k)]
    return sum(1 for g in grams if g in flat) / max(1, len(grams))


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 1
    folder = sys.argv[1]
    avoid = sys.argv[2:]
    tmp = tempfile.mkdtemp(prefix="exec_review_")
    try:
        print("转换中（WPS COM，隐藏窗口）...")
        made = convert_all(folder, tmp)
        print("%-36s %-4s %-16s %s" % ("文件", "页数", "落款", "提示"))
        print("-" * 100)
        for name, docx, pdf in made:
            pages = page_texts(pdf)
            joined = "\n".join(pages)
            flat = norm(joined)
            notes = []
            p_cizhi = next((i for i, t in enumerate(pages) if "此致" in t), None)
            if "此致" in flat:
                if p_cizhi is None:
                    notes.append("[!] 此致异常")
                else:
                    same = ("申请执行人" in pages[p_cizhi]) and ("年" in pages[p_cizhi])
                    if same:
                        notes.append("落款第%d页(整块同页)" % (p_cizhi + 1))
                    else:
                        notes.append("[!] 落款被拆页")
                    last_is_sign = pages[p_cizhi].strip().startswith("此致")
                    notes.append("末页只有落款" if last_is_sign else "落款与正文同页")
                    if name.startswith("01_") and len(pages) > 1:
                        if p_cizhi != len(pages) - 1:
                            notes.append("[!] 01落款未在末页")
                        elif not last_is_sign:
                            notes.append("[!] 01落款未独立成末页")
            para = [t for t in doc_texts(docx) if len(norm(t)) >= 30]
            low = sorted((shingle_ratio(t, flat), t) for t in para)
            hard = [(r, t) for r, t in low if r < 0.5]
            soft = [(r, t) for r, t in low if 0.5 <= r < 0.8]
            if hard:
                notes.append("[!] 疑似被裁行或渲染丢失 %d 段(命中率%.2f)，例如 %s"
                             % (len(hard), hard[0][0], hard[0][1][:20]))
            elif soft:
                notes.append("[?] 需人工确认 %d 段(命中率%.2f)，例如 %s"
                             % (len(soft), soft[0][0], soft[0][1][:20]))
            hit = [a for a in avoid if a and a in joined]
            if hit:
                notes.append("[!] 旧案残留：" + "、".join(hit))
            print("%-36s %-4d %-16s %s" % (name[:34], len(pages),
                                           ("第%d页" % (p_cizhi + 1)) if p_cizhi is not None else "-",
                                           "；".join(notes) or "OK"))
        print("-" * 100)
        # 交付前自检（字体/颜色/编号/串案/空段）一并跑掉，用户不用再敲第二条命令
        co = os.path.join(os.path.dirname(os.path.abspath(__file__)), "check_outputs.py")
        if os.path.isfile(co):
            print("\n=== 交付前自检（字体字号／颜色高亮／编号／串案／落款空段） ===")
            try:
                subprocess.run([sys.executable, co, folder] + list(avoid), check=False)
            except Exception as exc:
                print("  自检脚本运行失败：%s" % exc)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
